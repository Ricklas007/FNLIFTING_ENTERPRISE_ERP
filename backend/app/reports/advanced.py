import csv
import io
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Asset, Certificate, Client, Employee, Expense, Inspection, Invoice, PerformanceReview, Revenue, Trip


REPORT_TITLES = {
    "executive": "Relatório executivo consolidado",
    "financial": "Relatório financeiro avançado",
    "operations": "Relatório de operações técnicas",
    "hr": "Relatório de pessoas e desempenho",
    "fleet": "Relatório de frota e logística",
    "assets": "Relatório de património",
}


def _float(value) -> float:
    return float(value or 0)


def _money(value: float) -> str:
    return f"Kz {value:,.2f}".replace(",", " ")


def _months(start: date, end: date) -> list[str]:
    cursor = date(start.year, start.month, 1)
    result = []
    while cursor <= end and len(result) < 24:
        result.append(cursor.strftime("%Y-%m"))
        cursor = date(cursor.year + (cursor.month == 12), 1 if cursor.month == 12 else cursor.month + 1, 1)
    return result


def _group(rows, key, value=lambda _row: 1) -> list[dict]:
    totals: dict[str, float] = defaultdict(float)
    for row in rows:
        totals[str(key(row) or "Não definido")] += float(value(row) or 0)
    return [{"name": name, "value": round(amount, 2)} for name, amount in sorted(totals.items(), key=lambda item: item[1], reverse=True)]


def build_report(db: Session, company_id: str, report_type: str, date_from: date, date_to: date,
                 client_id: str | None = None, department_id: str | None = None) -> dict:
    expenses = list(db.scalars(select(Expense).where(
        Expense.company_id == company_id, Expense.deleted_at.is_(None),
        Expense.transaction_date.between(date_from, date_to),
        *([Expense.department_id == department_id] if department_id else []),
    )))
    revenues = list(db.scalars(select(Revenue).where(
        Revenue.company_id == company_id, Revenue.deleted_at.is_(None),
        Revenue.transaction_date.between(date_from, date_to),
        *([Revenue.client_id == client_id] if client_id else []),
    )))
    invoices = list(db.scalars(select(Invoice).where(
        Invoice.company_id == company_id, Invoice.deleted_at.is_(None),
        Invoice.issue_date.between(date_from, date_to),
        *([Invoice.client_id == client_id] if client_id else []),
    )))
    inspections = list(db.scalars(select(Inspection).where(
        Inspection.company_id == company_id, Inspection.deleted_at.is_(None),
        Inspection.inspection_date.between(date_from, date_to),
        *([Inspection.client_id == client_id] if client_id else []),
    )))
    trips = list(db.scalars(select(Trip).where(
        Trip.company_id == company_id, Trip.deleted_at.is_(None),
        Trip.created_at >= datetime.combine(date_from, datetime.min.time()),
        Trip.created_at <= datetime.combine(date_to, datetime.max.time()),
        *([Trip.client_id == client_id] if client_id else []),
    )))
    certificates = list(db.scalars(select(Certificate).where(
        Certificate.company_id == company_id, Certificate.deleted_at.is_(None),
        Certificate.issued_at.between(date_from, date_to),
        *([Certificate.client_id == client_id] if client_id else []),
    )))
    employees = list(db.scalars(select(Employee).where(
        Employee.company_id == company_id, Employee.deleted_at.is_(None),
        *([Employee.department_id == department_id] if department_id else []),
    )))
    assets = list(db.scalars(select(Asset).where(Asset.company_id == company_id, Asset.deleted_at.is_(None))))
    performances = list(db.scalars(select(PerformanceReview).where(
        PerformanceReview.company_id == company_id, PerformanceReview.deleted_at.is_(None),
        PerformanceReview.period >= date_from.strftime("%Y-%m"), PerformanceReview.period <= date_to.strftime("%Y-%m"),
    )))
    clients = {item.id: item.name for item in db.scalars(select(Client).where(Client.company_id == company_id))}

    revenue_total = sum(_float(row.amount) for row in revenues)
    expense_total = sum(_float(row.amount) for row in expenses)
    invoiced_total = sum(_float(row.total_amount) for row in invoices)
    paid_total = sum(_float(payment.amount) for invoice in invoices for payment in invoice.payments if payment.status == "confirmed")
    receivable = max(0, invoiced_total - paid_total)
    average_performance = sum(_float(item.overall_score) for item in performances) / len(performances) if performances else 0
    asset_value = sum(_float(item.current_value) for item in assets)
    metrics = [
        {"key": "revenue", "label": "Receita", "value": revenue_total, "display_value": _money(revenue_total)},
        {"key": "expense", "label": "Despesa", "value": expense_total, "display_value": _money(expense_total)},
        {"key": "profit", "label": "Resultado", "value": revenue_total - expense_total, "display_value": _money(revenue_total - expense_total)},
        {"key": "receivable", "label": "A receber", "value": receivable, "display_value": _money(receivable)},
        {"key": "inspections", "label": "Inspeções", "value": len(inspections), "display_value": str(len(inspections))},
        {"key": "certificates", "label": "Certificados", "value": len(certificates), "display_value": str(len(certificates))},
        {"key": "trips", "label": "Viagens", "value": len(trips), "display_value": str(len(trips))},
        {"key": "employees", "label": "Trabalhadores ativos", "value": sum(row.status == "active" for row in employees), "display_value": str(sum(row.status == "active" for row in employees))},
        {"key": "performance", "label": "Desempenho médio", "value": average_performance, "display_value": f"{average_performance:.1f}%"},
        {"key": "assets", "label": "Valor patrimonial", "value": asset_value, "display_value": _money(asset_value)},
    ]

    month_rows = {month: {"period": month, "revenue": 0.0, "expense": 0.0, "inspections": 0, "trips": 0} for month in _months(date_from, date_to)}
    for row in revenues:
        month_rows.setdefault(row.transaction_date.strftime("%Y-%m"), {"period": row.transaction_date.strftime("%Y-%m"), "revenue": 0, "expense": 0, "inspections": 0, "trips": 0})["revenue"] += _float(row.amount)
    for row in expenses:
        month_rows.setdefault(row.transaction_date.strftime("%Y-%m"), {"period": row.transaction_date.strftime("%Y-%m"), "revenue": 0, "expense": 0, "inspections": 0, "trips": 0})["expense"] += _float(row.amount)
    for row in inspections:
        month_rows[row.inspection_date.strftime("%Y-%m")]["inspections"] += 1
    for row in trips:
        month_rows[row.created_at.strftime("%Y-%m")]["trips"] += 1

    highlights = []
    if revenue_total - expense_total >= 0:
        highlights.append(f"Resultado positivo de {_money(revenue_total - expense_total)} no período.")
    else:
        highlights.append(f"Resultado negativo de {_money(abs(revenue_total - expense_total))}; requer análise de custos.")
    if inspections:
        approved = sum(row.status == "approved" for row in inspections)
        highlights.append(f"{approved} de {len(inspections)} inspeções foram aprovadas ({approved / len(inspections) * 100:.0f}%).")
    if receivable:
        highlights.append(f"Contas a receber em aberto: {_money(receivable)}.")
    if average_performance:
        highlights.append(f"Desempenho médio dos trabalhadores: {average_performance:.1f}%.")

    return {
        "report_type": report_type,
        "title": REPORT_TITLES.get(report_type, REPORT_TITLES["executive"]),
        "date_from": date_from,
        "date_to": date_to,
        "generated_at": datetime.now(timezone.utc),
        "metrics": metrics,
        "series": list(month_rows.values()),
        "breakdowns": {
            "Despesas por categoria": _group(expenses, lambda row: row.category, lambda row: row.amount),
            "Receitas por cliente": _group(revenues, lambda row: clients.get(row.client_id, row.client_id), lambda row: row.amount),
            "Inspeções por estado": _group(inspections, lambda row: row.status),
            "Viagens por província": _group(trips, lambda row: row.province),
            "Património por categoria": _group(assets, lambda row: row.category, lambda row: row.current_value),
        },
        "highlights": highlights,
    }


def export_pdf(report: dict) -> bytes:
    output = io.BytesIO()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title", parent=styles["Title"], textColor=colors.HexColor("#0b56d8"), alignment=TA_CENTER, fontSize=18, spaceAfter=8)
    document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    story = [Paragraph("FNLIFTING ENTERPRISE ERP", title_style), Paragraph(report["title"], styles["Heading2"]), Paragraph(f"Período: {report['date_from']} a {report['date_to']}", styles["Normal"]), Spacer(1, 5 * mm)]
    metric_rows = [["Indicador", "Valor"]] + [[item["label"], item["display_value"]] for item in report["metrics"]]
    metric_table = Table(metric_rows, colWidths=[85 * mm, 45 * mm], repeatRows=1)
    metric_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b56d8")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#b6c8e8")), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f6ff")]), ("PADDING", (0, 0), (-1, -1), 6)]))
    story.extend([metric_table, Spacer(1, 5 * mm), Paragraph("Destaques", styles["Heading2"])])
    for highlight in report["highlights"]:
        story.append(Paragraph(f"• {highlight}", styles["BodyText"]))
    story.append(Spacer(1, 4 * mm))
    for name, rows in report["breakdowns"].items():
        if not rows:
            continue
        story.append(Paragraph(name, styles["Heading3"]))
        table = Table([["Item", "Valor"]] + [[row["name"], f"{row['value']:,.2f}"] for row in rows[:12]], colWidths=[95 * mm, 45 * mm], repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#13294b")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), .3, colors.HexColor("#cbd5e1")), ("PADDING", (0, 0), (-1, -1), 5)]))
        story.extend([table, Spacer(1, 3 * mm)])
    document.build(story)
    return output.getvalue()


def export_xlsx(report: dict) -> bytes:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Resumo"
    overview.append(["FNLIFTING ENTERPRISE ERP", report["title"]])
    overview.append(["Período", f"{report['date_from']} a {report['date_to']}"])
    overview.append([])
    overview.append(["Indicador", "Valor", "Valor numérico"])
    for item in report["metrics"]:
        overview.append([item["label"], item["display_value"], item["value"]])
    monthly = workbook.create_sheet("Evolução mensal")
    monthly.append(["Período", "Receitas", "Despesas", "Inspeções", "Viagens"])
    for row in report["series"]:
        monthly.append([row["period"], row["revenue"], row["expense"], row["inspections"], row["trips"]])
    details = workbook.create_sheet("Detalhes")
    details.append(["Grupo", "Item", "Valor"])
    for group, rows in report["breakdowns"].items():
        for row in rows:
            details.append([group, row["name"], row["value"]])
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0B56D8")
        for column in sheet.columns:
            width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
        sheet.freeze_panes = "A2"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_csv(report: dict) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(["FNLIFTING ENTERPRISE ERP", report["title"]])
    writer.writerow(["Período", report["date_from"], report["date_to"]])
    writer.writerow([])
    writer.writerow(["Indicador", "Valor", "Valor numérico"])
    for item in report["metrics"]:
        writer.writerow([item["label"], item["display_value"], item["value"]])
    writer.writerow([])
    writer.writerow(["Grupo", "Item", "Valor"])
    for group, rows in report["breakdowns"].items():
        for row in rows:
            writer.writerow([group, row["name"], row["value"]])
    return ("\ufeff" + output.getvalue()).encode("utf-8")
